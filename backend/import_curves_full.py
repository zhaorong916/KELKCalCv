"""导入完整曲线数据: Eq% Data / FL Data / Kc Data / Fd Data

从 KELK CalcCv Valve & Trim Factors.xlsm 导入:
  - Eq% Data: A0~A6 系数 (eqcv), 索引 1~198
  - FL Data: P0~P4 + Min/Max (fl), 索引 1~7
  - Kc Data: P0~P4 + Min/Max (kc), 索引 1~7
  - Fd Data: P0~P4 + Min/Max (fd), 索引 1~7

同时导入每个系列 Sheet 中 "Choices for Graphs" 区域的
Plug Size → %Stk Index 映射表 (按压力等级分组)。
"""
import os
import openpyxl
from sqlalchemy import text
from app.core.database import engine

EXCEL_PATH = os.path.join(
    os.path.dirname(__file__), '..', '参考文件',
    'KELK CalcCv Valve & Trim Factors.xlsm'
)

# 各系列 Sheet 中 Plug Size → %Stk Index 查找表的起始行
# (GC120 中在 B80:D105 区域，其他系列类似)
PLUG_INDEX_TABLE_START = {
    'GC120': 80,
    'GC210': 80,
    # 其他系列后续补充
}


def import_eq_curves(wb, conn):
    """导入 Eq% Data sheet 的所有 A0~A6 系数"""
    ws = wb['Eq% Data']
    count = 0
    for r in range(6, ws.max_row + 1):
        idx = ws.cell(r, 1).value
        if idx is None:
            continue
        try:
            idx = int(idx)
        except (ValueError, TypeError):
            continue
        a6 = ws.cell(r, 2).value
        a5 = ws.cell(r, 3).value
        a4 = ws.cell(r, 4).value
        a3 = ws.cell(r, 5).value
        a2 = ws.cell(r, 6).value
        a1 = ws.cell(r, 7).value
        a0 = ws.cell(r, 8).value
        if a0 is None and a6 is None:
            continue
        conn.execute(text("""
            MERGE valve_factor_curves AS target
            USING (SELECT 'eqcv' AS factor_type, :idx AS eq_index) AS source
            ON target.factor_type = source.factor_type AND target.eq_index = source.eq_index
            WHEN MATCHED THEN
                UPDATE SET a0=:a0, a1=:a1, a2=:a2, a3=:a3, a4=:a4, a5=:a5, a6=:a6
            WHEN NOT MATCHED THEN
                INSERT (factor_type, eq_index, a0, a1, a2, a3, a4, a5, a6)
                VALUES ('eqcv', :idx, :a0, :a1, :a2, :a3, :a4, :a5, :a6);
        """), {'idx': idx, 'a0': a0, 'a1': a1, 'a2': a2, 'a3': a3, 'a4': a4, 'a5': a5, 'a6': a6})
        count += 1
    conn.commit()
    print(f"  Eq% 曲线导入: {count} 条")
    return count


def import_factor_curves(wb, sheet_name, factor_type, conn):
    """导入 FL/Kc/Fd Data sheet 的 P0~P4 + Min/Max"""
    ws = wb[sheet_name]
    count = 0
    for r in range(6, ws.max_row + 1):
        idx = ws.cell(r, 1).value
        if idx is None:
            continue
        try:
            idx = int(idx)
        except (ValueError, TypeError):
            continue
        vmin = ws.cell(r, 2).value
        vmax = ws.cell(r, 3).value
        p4 = ws.cell(r, 4).value
        p3 = ws.cell(r, 5).value
        p2 = ws.cell(r, 6).value
        p1 = ws.cell(r, 7).value
        p0 = ws.cell(r, 8).value
        if p0 is None and p4 is None and vmin is None:
            continue
        conn.execute(text("""
            MERGE valve_factor_curves AS target
            USING (SELECT :ft AS factor_type, :idx AS eq_index) AS source
            ON target.factor_type = source.factor_type AND target.eq_index = source.eq_index
            WHEN MATCHED THEN
                UPDATE SET value_min=:vmin, value_max=:vmax,
                           p0=:p0, p1=:p1, p2=:p2, p3=:p3, p4=:p4
            WHEN NOT MATCHED THEN
                INSERT (factor_type, eq_index, value_min, value_max, p0, p1, p2, p3, p4)
                VALUES (:ft, :idx, :vmin, :vmax, :p0, :p1, :p2, :p3, :p4);
        """), {'ft': factor_type, 'idx': idx, 'vmin': vmin, 'vmax': vmax,
              'p0': p0, 'p1': p1, 'p2': p2, 'p3': p3, 'p4': p4})
        count += 1
    conn.commit()
    print(f"  {factor_type.upper()} 曲线导入: {count} 条")
    return count


def import_plug_index_mapping(wb, conn):
    """导入各系列 Sheet 中 Plug Size → %Stk Index 映射表

    表结构 (GC120 为例):
      B80: ASME Pressure Rating | C80: Plug Size | D80: %Stk Index
      B83: 900# /1500# | C83: 18 | D83: 1
      ...
    """
    # 先确保表存在
    conn.execute(text("""
        IF NOT EXISTS (SELECT * FROM sysobjects WHERE name='valve_plug_index_map' AND xtype='U')
        CREATE TABLE valve_plug_index_map (
            id INT IDENTITY(1,1) PRIMARY KEY,
            series_code VARCHAR(32) NOT NULL,
            asme_class VARCHAR(32) NOT NULL,
            plug_mm FLOAT NOT NULL,
            stk_index INT NOT NULL,
            CONSTRAINT uq_plug_map UNIQUE (series_code, asme_class, plug_mm)
        )
    """))
    conn.commit()

    # 清理旧数据
    conn.execute(text("DELETE FROM valve_plug_index_map"))
    conn.commit()

    series_sheets = ['GC120', 'GC210', 'GV120', 'GV110', 'GV410',
                     'BC110', 'BE120', 'BT120', 'BT110', 'BT210']
    total = 0

    for sn in series_sheets:
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        # 查找映射表起始行 (包含 "Plug Size" 和 "%Stk Index" 的行)
        start_row = None
        for r in range(1, min(ws.max_row + 1, 150)):
            c_val = ws.cell(r, 3).value
            d_val = ws.cell(r, 4).value
            if c_val and 'Plug' in str(c_val) and d_val and 'Stk' in str(d_val):
                start_row = r + 2  # 表头+空行后开始
                break
        if start_row is None:
            continue

        current_asme = None
        for r in range(start_row, min(ws.max_row + 1, start_row + 60)):
            b_val = ws.cell(r, 2).value
            c_val = ws.cell(r, 3).value
            d_val = ws.cell(r, 4).value

            # B 列可能是新的压力等级
            if b_val and str(b_val).strip() and '#' in str(b_val):
                current_asme = str(b_val).strip()
                continue

            if c_val is None or d_val is None:
                continue
            try:
                plug_mm = float(c_val)
                stk_idx = int(d_val)
            except (ValueError, TypeError):
                continue

            if current_asme is None:
                continue

            conn.execute(text("""
                INSERT INTO valve_plug_index_map (series_code, asme_class, plug_mm, stk_index)
                VALUES (:sn, :asme, :plug, :idx)
            """), {'sn': sn, 'asme': current_asme, 'plug': plug_mm, 'idx': stk_idx})
            total += 1

    conn.commit()
    print(f"  Plug Size→%Stk Index 映射导入: {total} 条")
    return total


def main():
    print("=" * 60)
    print("导入完整曲线数据")
    print("=" * 60)

    if not os.path.exists(EXCEL_PATH):
        print(f"[ERROR] Excel 文件不存在: {EXCEL_PATH}")
        return

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True, keep_vba=True)

    with engine.connect() as conn:
        print("\n[1] 导入 Eq% 曲线 (A0~A6)")
        import_eq_curves(wb, conn)

        print("\n[2] 导入 FL 曲线 (P0~P4)")
        import_factor_curves(wb, 'FL Data', 'fl', conn)

        print("\n[3] 导入 Kc 曲线 (P0~P4)")
        import_factor_curves(wb, 'Kc Data', 'kc', conn)

        print("\n[4] 导入 Fd 曲线 (P0~P4)")
        import_factor_curves(wb, 'Fd Data', 'fd', conn)

        print("\n[5] 导入 Plug Size → %Stk Index 映射")
        import_plug_index_mapping(wb, conn)

    wb.close()
    print("\n" + "=" * 60)
    print("导入完成")
    print("=" * 60)


if __name__ == '__main__':
    main()
