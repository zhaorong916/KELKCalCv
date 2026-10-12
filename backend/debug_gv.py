"""快速调试 GV120"""
import openpyxl

path = r'd:\KELK办公\信息化工作\207 自研系统\KELKCalCv\参考文件\KELK CalcCv Valve & Trim Factors.xlsm'
wb = openpyxl.load_workbook(path, data_only=True, keep_vba=True)
ws = wb['GV120']

# 模拟 globe_sizeonly 逻辑
header_r = None
for r in range(1, 20):
    a_val = ws.cell(row=r, column=1).value
    b_val = ws.cell(row=r, column=2).value
    if a_val and 'Trim Type' in str(a_val):
        header_r = r
        print(f'  header from A: r={r}, val={a_val}')
        break
    if b_val and 'Body and Trim Size' in str(b_val):
        header_r = r
        print(f'  header from B: r={r}, val={b_val}')
        break

print(f'header_r = {header_r}')
data_start = header_r + 3
print(f'data_start = {data_start}')

for r in range(data_start, data_start + 5):
    a = ws.cell(row=r, column=1).value
    b = ws.cell(row=r, column=2).value
    c = ws.cell(row=r, column=3).value
    d = ws.cell(row=r, column=4).value
    e = ws.cell(row=r, column=5).value
    f = ws.cell(row=r, column=6).value
    print(f'R{r}: A={repr(a)} B={repr(b)} C={repr(c)} D={repr(d)} E={repr(e)} F={repr(f)}')

wb.close()
