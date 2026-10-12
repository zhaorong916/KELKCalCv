import openpyxl

path = r'd:\KELK办公\信息化工作\207 自研系统\KELKCalCv\参考文件\KELK CalcCv Valve & Trim Factors.xlsm'
wb = openpyxl.load_workbook(path, data_only=True, keep_vba=True)

for sn in ['GC210', 'BT120', 'BT210', 'GV120', 'BC110', 'BE120']:
    if sn not in wb.sheetnames:
        continue
    ws = wb[sn]
    print(f'\n=== {sn} Sheet (A-G, R1-R30) ===')
    for r in range(1, 31):
        row_vals = []
        for c in range(1, 8):
            v = ws.cell(row=r, column=c).value
            row_vals.append(str(v) if v is not None else '')
        if any(row_vals):
            print(f'R{r:02d}: ' + ' | '.join(row_vals))

wb.close()
